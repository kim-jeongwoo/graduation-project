#!/usr/bin/env python3
"""Phase E validation: check parsing correctness, score distributions, ICC details."""
import os
import csv
import unicodedata
import docx
import numpy as np
from collections import defaultdict

VOCAL_DIR = '/home/dmlserver/projects/egg_prediction/data/vocal_academy'
EVAL_DIR  = os.path.join(VOCAL_DIR, '_new_upload/eval_sheets')
SCORE_COLS = ['비강사용', '흉성두성', '접촉강도', '호흡식별', '전방공명',
              '연구개위치', '공명개방', '혀경직도', '상압하압']

def nfc(s): return unicodedata.normalize('NFC', s)

def load_csv(path):
    with open(path, encoding='utf-8') as f:
        return list(csv.DictReader(f))

def find_docx(rater_n):
    fname = f'평가지{rater_n}.docx'
    path = os.path.join(EVAL_DIR, fname)
    if not os.path.exists(path):
        for f in os.listdir(EVAL_DIR):
            if nfc(f) == fname:
                return os.path.join(EVAL_DIR, f)
    return path

def get_item_rows_from_docx(rater_n):
    """Return list of (item_num, register, scores_list) from docx."""
    path = find_docx(rater_n)
    doc = docx.Document(path)
    rows = []
    for ti in range(2, len(doc.tables)):
        for row in doc.tables[ti].rows:
            cells = [nfc(c.text.strip()) for c in row.cells]
            if len(cells) < 12: continue
            item_str = cells[1]
            register = cells[2]
            if not item_str.isdigit(): continue
            scores = []
            for s in cells[3:12]:
                s = s.strip()
                scores.append(int(s) if s.lstrip('-').isdigit() else None)
            rows.append((int(item_str), register, scores))
    return rows

# ─────────────────────────────────────────────────────────────
# V-1: Item number consistency across 5 sheets
# ─────────────────────────────────────────────────────────────
print("="*70)
print("V-1: Item number consistency across 평가지1~5")
print("="*70)

all_item_sets = {}
for rater_n in range(1, 6):
    rows = get_item_rows_from_docx(rater_n)
    items = sorted(set(r[0] for r in rows))
    all_item_sets[rater_n] = items
    print(f"  평가지{rater_n}: items={items} (n={len(rows)} rows)")

# Check if all same
ref = all_item_sets[1]
for rn in range(2, 6):
    if all_item_sets[rn] != ref:
        print(f"  ⚠️ 평가지{rn} items differ from 평가지1!")

# ─────────────────────────────────────────────────────────────
# V-2: 평가지3 parsing correctness
# ─────────────────────────────────────────────────────────────
print("\n" + "="*70)
print("V-2: 평가지3 vs 평가지1 score comparison (same items)")
print("="*70)

rows1 = {(r, reg): s for r, reg, s in get_item_rows_from_docx(1)}
rows3 = {(r, reg): s for r, reg, s in get_item_rows_from_docx(3)}

print(f"  평가지1: {len(rows1)} (item, register) combos")
print(f"  평가지3: {len(rows3)} (item, register) combos")

# Show item 1 저/중/고 for both
print("\n  Item 1 scores (평가지1 vs 평가지3):")
for reg in ['저', '중', '고']:
    s1 = rows1.get((1, reg), ['?']*9)
    s3 = rows3.get((1, reg), ['?']*9)
    print(f"    {reg}: 평가지1={s1} | 평가지3={s3}")

# ─────────────────────────────────────────────────────────────
# V-3: Score distribution sanity check
# ─────────────────────────────────────────────────────────────
print("\n" + "="*70)
print("V-3: Score distribution by singer group and per rater")
print("="*70)

v2  = load_csv(os.path.join(VOCAL_DIR, 'trainer_scores_v2.csv'))
full = load_csv(os.path.join(VOCAL_DIR, 'trainer_scores.csv'))

# Stage1 vs Stage2 score distribution
def score_dist(rows, col='비강사용'):
    vals = [int(r[col]) for r in rows if r.get(col, '').isdigit()]
    if not vals: return {}
    counts = {i: vals.count(i) for i in range(6)}
    return counts

print("\n  비강사용 점수 분포:")
stage1_rows = [r for r in full if int(r.get('singer', 0)) <= 12]
stage2_rows = [r for r in full if int(r.get('singer', 0)) > 12]
d1 = score_dist(stage1_rows)
d2 = score_dist(stage2_rows)
print(f"  Singer 1~12:  {d1}")
print(f"  Singer 13~22: {d2}")

# Per-rater average
print("\n  Per-rater mean score (전방공명, full dataset):")
for rater in range(1, 6):
    rater_rows = [r for r in full if r.get('rater', '') == str(rater)]
    vals = [float(r.get('전방공명', '0')) for r in rater_rows if r.get('전방공명', '').isdigit()]
    if vals:
        print(f"    Rater {rater}: mean={np.mean(vals):.2f} std={np.std(vals):.2f} n={len(vals)}")

# ─────────────────────────────────────────────────────────────
# V-4: ICC details — per item, per rater variance
# ─────────────────────────────────────────────────────────────
print("\n" + "="*70)
print("V-4: ICC detailed analysis + outlier raters")
print("="*70)

def icc_21(data):
    subjects = sorted(set(s for s, _, _ in data))
    raters   = sorted(set(r for _, r, _ in data))
    n, k = len(subjects), len(raters)
    if n < 2 or k < 2: return None
    sub_idx = {s: i for i, s in enumerate(subjects)}
    mat = np.full((n, k), np.nan)
    for s, r, v in data:
        ri = raters.index(r)
        mat[sub_idx[s]][ri] = float(v)
    row_m = np.nanmean(mat, axis=1)
    col_m = np.nanmean(mat, axis=0)
    grand = np.nanmean(mat)
    SS_r = k * np.nansum((row_m - grand)**2)
    SS_c = n * np.nansum((col_m - grand)**2)
    SS_tot = np.nansum((mat - grand)**2)
    SS_e = SS_tot - SS_r - SS_c
    df_r, df_c = n-1, k-1
    df_e = (n-1)*(k-1)
    MS_r = SS_r/df_r; MS_c = SS_c/df_c; MS_e = SS_e/df_e if df_e > 0 else 1e-10
    icc = (MS_r - MS_e) / (MS_r + (k-1)*MS_e + k*(MS_c - MS_e)/n)
    return float(icc), mat, subjects, raters

# Stage2 full: per-item ICC and outlier rater
print("\n  Stage2 per-item ICC + outlier rater detection:")
for col in SCORE_COLS:
    data = []
    for row in full:
        singer = row.get('singer', ''); reg = row.get('register', '')
        rater = row.get('rater', ''); val = row.get(col, '')
        if val.isdigit():
            data.append((f"{singer}_{reg}", rater, int(val)))
    result = icc_21(data)
    if result is None:
        print(f"  {col}: N/A")
        continue
    icc_val, mat, subjects, raters = result
    # Per-rater mean deviation from group mean
    rater_devs = []
    for ri, rat in enumerate(raters):
        col_vals = mat[:, ri]
        valid = col_vals[~np.isnan(col_vals)]
        others = mat.copy(); others[:, ri] = np.nan
        other_means = np.nanmean(others, axis=1)
        valid_other = other_means[~np.isnan(col_vals)]
        dev = np.mean(np.abs(valid - valid_other)) if len(valid) > 0 else 0
        rater_devs.append((rat, dev))
    rater_devs.sort(key=lambda x: -x[1])
    outlier = rater_devs[0]
    print(f"  {col:6s}: ICC={icc_val:.3f}  "
          f"outlier=rater{outlier[0]} (dev={outlier[1]:.2f})")

# ─────────────────────────────────────────────────────────────
# V-5: Validation CSV consistency
# ─────────────────────────────────────────────────────────────
print("\n" + "="*70)
print("V-5: Validation CSV consistency check")
print("="*70)

val = load_csv(os.path.join(VOCAL_DIR, 'trainer_scores_validation.csv'))
print(f"  Total rows: {len(val)} (expected 165 = 11×5×3)")

# Per test_id row count
test_counts = defaultdict(int)
for row in val:
    test_counts[row['test_id']] += 1
print("\n  Rows per test_id:")
for tid in sorted(test_counts):
    n = test_counts[tid]
    ok = '✅' if n == 15 else f'⚠️ n={n}'
    print(f"    {tid}: {n} {ok}")

# Score means per test_id (전방공명)
print("\n  전방공명 mean per test_id:")
test_scores = defaultdict(list)
for row in val:
    v = row.get('전방공명', '')
    if v.isdigit():
        test_scores[row['test_id']].append(int(v))
for tid in sorted(test_scores):
    vals = test_scores[tid]
    print(f"    {tid}: mean={np.mean(vals):.2f} std={np.std(vals):.2f}")
