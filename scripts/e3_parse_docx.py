#!/usr/bin/env python3
"""
E-3: Parse 평가지1~5.docx into trainer_scores_v2.csv and trainer_scores_validation.csv.

Structure per docx:
  Table 0: rater ID + date
  Table 1: column headers (ignore)
  Table 2: items 1~7  (rows: item_num, register, ①~⑨)
  Table 3: items 8~14
  Table 4: items 15~21

item_num (in '곡 번호' col): 1~10 → singer_13~22 (training)
                             11~21 → test_01~11  (validation)
"""
import os
import csv
import unicodedata
import docx

EVAL_DIR = '/home/dmlserver/projects/egg_prediction/data/vocal_academy/_new_upload/eval_sheets'
OUT_DIR = '/home/dmlserver/projects/egg_prediction/data/vocal_academy'

SINGER_OFFSET = 12      # item N (1~10) → singer (12+N)
SCORE_COLS = ['비강사용', '흉성두성', '접촉강도', '호흡식별', '전방공명', '연구개위치', '공명개방', '혀경직도', '상압하압']

# Singer gender for new singers 13~22 (folder 1~10)
GENDER_NEW = {13:'M', 14:'M', 15:'F', 16:'M', 17:'M',
              18:'M', 19:'M', 20:'F', 21:'M', 22:'M'}
# Test gender for test_01~11
GENDER_TEST = {1:'M', 2:'M', 3:'F', 4:'M', 5:'F',
               6:'M', 7:'M', 8:'F', 9:'M', 10:'M', 11:'M'}

# Song mapping for training singers: item_num → song_id
# From E-1: folder1→song4, 2→4, 3→2, 4→6, 5→6, 6→6, 7→14, 8→8, 9→4, 10→4
ITEM_SONG = {1:4, 2:4, 3:2, 4:6, 5:6, 6:6, 7:14, 8:8, 9:4, 10:4}

SONG_NAMES = {
    2: '윤하 - 기다리다',
    4: '김범수 - 보고싶다',
    6: '폴킴 - 모든 날, 모든 순간',
    8: '이하이 - 한숨',
    14: '성시경 - 너의 모든 순간',
}

REG_MAP = {'저': '저', '중': '중', '고': '고'}

def nfc(s):
    return unicodedata.normalize('NFC', s)

def extract_tables(doc_path):
    """Extract all data rows from tables in a docx."""
    doc = docx.Document(doc_path)
    # Table 0 → rater ID (fallback to filename index if cell empty)
    t0 = doc.tables[0]
    rater_cell = nfc(t0.rows[0].cells[1].text.strip())
    if rater_cell.isdigit():
        rater_id = int(rater_cell)
    else:
        # Infer from filename (평가지N.docx → N)
        import re as _re
        m = _re.search(r'평가지(\d+)', doc_path)
        rater_id = int(m.group(1)) if m else 0
        print(f"  ⚠️ Empty rater ID in {os.path.basename(doc_path)}, inferred as {rater_id}")
    
    rows = []
    # Tables from index 2 onwards contain data (Table 0=meta, Table 1=legend)
    for ti in range(2, len(doc.tables)):
        table = doc.tables[ti]
        for row in table.rows:  # may or may not have header row
            cells = [nfc(c.text.strip()) for c in row.cells]
            if len(cells) < 14:
                continue
            item_num_str = cells[1]
            register = cells[2]
            scores = cells[3:12]  # ①~⑨
            if not item_num_str.isdigit():  # skip header rows
                continue
            item_num = int(item_num_str)
            if not register:
                continue
            score_vals = []
            for s in scores:
                s = s.strip()
                if s == '' or not s.lstrip('-').isdigit():
                    score_vals.append(None)
                else:
                    score_vals.append(int(s))
            rows.append({
                'rater': rater_id,
                'item_num': item_num,
                'register': register,
                'scores': score_vals,
            })
    return rater_id, rows


def parse_all_docx():
    training_rows = []
    validation_rows = []
    warnings = []

    for rater in range(1, 6):
        fname = f'평가지{rater}.docx'
        path = os.path.join(EVAL_DIR, fname)
        # Try NFC/NFD for macOS filenames
        if not os.path.exists(path):
            for f in os.listdir(EVAL_DIR):
                if nfc(f) == fname:
                    path = os.path.join(EVAL_DIR, f)
                    break

        if not os.path.exists(path):
            warnings.append(f"Missing: {fname}")
            continue

        rater_id, rows = extract_tables(path)
        print(f"평가지{rater}: rater_id={rater_id}, {len(rows)} data rows")

        for r in rows:
            item_num = r['item_num']
            register = r['register']
            scores = r['scores']

            if item_num <= 10:
                # Training: singer_13~22
                singer_id = SINGER_OFFSET + item_num
                song_id = ITEM_SONG.get(item_num)
                song_name = SONG_NAMES.get(song_id, '')
                gender = GENDER_NEW.get(singer_id, '')
                training_rows.append({
                    'rater': rater_id, 'singer': singer_id, 'gender': gender,
                    'song': song_id, 'song_name': song_name, 'register': register,
                    **{SCORE_COLS[i]: scores[i] for i in range(9)},
                })
            else:
                # Validation: test_01~11
                test_id = item_num - 10  # item 11→test 1, ..., 21→test 11
                gender = GENDER_TEST.get(test_id, '')
                validation_rows.append({
                    'rater': rater_id, 'test_id': f'test_{test_id:02d}', 'gender': gender,
                    'register': register,
                    **{SCORE_COLS[i]: scores[i] for i in range(9)},
                })

    return training_rows, validation_rows, warnings


def write_csv(rows, path, fieldnames):
    with open(path, 'w', encoding='utf-8', newline='\n') as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def main():
    training_rows, validation_rows, warnings = parse_all_docx()

    # Apply manual corrections: item 13 저음역 ⑨ 6→5
    # (평가지1 + 평가지4 both had 6 in that cell)
    corrected = 0
    for row in validation_rows:
        if row['test_id'] == 'test_03' and row['register'] == '저':
            if row.get('상압하압') == 6:
                row['상압하압'] = 5
                corrected += 1
                print(f"Applied correction: rater={row['rater']} test_03 저 상압하압 6→5")

    # Post-correction range check
    for row in training_rows + validation_rows:
        subject = row.get('singer') or row.get('test_id')
        for col in SCORE_COLS:
            s = row.get(col)
            if s is not None and (s < 0 or s > 5):
                warnings.append(f"Out of range: rater={row['rater']} {subject} {row['register']} {col}={s}")

    train_fields = ['rater', 'singer', 'gender', 'song', 'song_name', 'register'] + SCORE_COLS
    val_fields = ['rater', 'test_id', 'gender', 'register'] + SCORE_COLS

    train_path = os.path.join(OUT_DIR, 'trainer_scores_v2.csv')
    val_path = os.path.join(OUT_DIR, 'trainer_scores_validation.csv')

    write_csv(training_rows, train_path, train_fields)
    write_csv(validation_rows, val_path, val_fields)

    print(f"\n✅ trainer_scores_v2.csv: {len(training_rows)} rows → {train_path}")
    print(f"✅ trainer_scores_validation.csv: {len(validation_rows)} rows → {val_path}")

    if warnings:
        print(f"\n⚠️ Warnings ({len(warnings)}):")
        for w in warnings:
            print(f"  {w}")
    else:
        print("\n✅ No warnings (all scores in 0~5 range, no missing files)")

    # Spot check
    print("\nSample training rows (first 3):")
    for row in training_rows[:3]:
        print(f"  {row}")
    print("\nSample validation rows (first 3):")
    for row in validation_rows[:3]:
        print(f"  {row}")


if __name__ == '__main__':
    main()
